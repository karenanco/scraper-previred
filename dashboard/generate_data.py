import json
import os
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

EXCEL_PATH = Path(__file__).parent.parent / 'data' / 'indicadores.xlsx'
DATA_JSON_PATH = Path(__file__).parent / 'data.json'


def _try_parse_json(val):
    if val is None:
        return None
    if isinstance(val, (dict, list)):
        return val
    if isinstance(val, str):
        try:
            return json.loads(val)
        except (json.JSONDecodeError, ValueError):
            return val
    if isinstance(val, float):
        return val
    return val


def generate():
    if not EXCEL_PATH.exists():
        print("No se encontró el archivo de datos en:", EXCEL_PATH)
        return

    wb = load_workbook(EXCEL_PATH)
    result = {}

    for sheet_name, key in [('Previred', 'previred'), ('UTM_UTA', 'utm')]:
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in row):
                row_data = {}
                for i, v in enumerate(row):
                    row_data[headers[i]] = _try_parse_json(v)
                rows.append(row_data)
        result[key] = {'fields': headers, 'rows': rows}

    mtime = os.path.getmtime(EXCEL_PATH)
    result['last_updated'] = datetime.fromtimestamp(mtime).strftime('%d-%m-%Y %H:%M')

    with open(DATA_JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"Datos generados correctamente: {DATA_JSON_PATH}")


if __name__ == '__main__':
    generate()
