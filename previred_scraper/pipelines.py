import json
import os
from openpyxl import Workbook, load_workbook

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data')
EXCEL_PATH = os.path.join(DATA_DIR, 'indicadores.xlsx')

PREVIRED_FIELDS = [
    'mes', 'periodo_remuneracion',
    'uf_valor', 'uf_fecha',
    'renta_tope_afp', 'renta_tope_ips', 'renta_tope_cesantia',
    'afp_tasas', 'afc_tasas',
    'apv_tope_mensual', 'apv_tope_anual', 'deposito_convenido_tope',
    'utm_valor', 'uta_valor',
    'renta_minima_dependientes', 'renta_minima_menores',
    'renta_minima_casa_particular', 'renta_minima_no_remuneracional',
    'seguro_social_tasa', 'sis_tasa',
    'salud_ccaf', 'salud_fonasa',
    'trabajo_pesado_tasa', 'trabajo_menos_pesado_tasa',
    'asignacion_familiar',
]

UTM_FIELDS = [
    'mes', 'utm', 'uta', 'ipc_puntos',
    'variacion_mensual', 'variacion_acumulada', 'variacion_anual',
]


def _init_excel():
    os.makedirs(DATA_DIR, exist_ok=True)
    if os.path.exists(EXCEL_PATH):
        return
    wb = Workbook()
    ws_prev = wb.active
    ws_prev.title = 'Previred'
    ws_prev.append(PREVIRED_FIELDS)
    ws_utm = wb.create_sheet('UTM_UTA')
    ws_utm.append(UTM_FIELDS)
    wb.save(EXCEL_PATH)


def _to_excel_val(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _upsert_rows(rows, key_field, fields):
    _init_excel()
    wb = load_workbook(EXCEL_PATH)

    sheet_name = 'Previred' if fields == PREVIRED_FIELDS else 'UTM_UTA'
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
        ws.append(fields)
    else:
        ws = wb[sheet_name]

    existing = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            existing[str(row[0]).strip().lower()] = row

    header = [c.value for c in ws[1]]
    key_idx = header.index(key_field)

    for row_data in rows:
        key = str(row_data.get(key_field, '')).strip().lower()
        if key in existing:
            existing_row_data = existing[key]
            for col_idx, field in enumerate(fields):
                if field in row_data and row_data[field] is not None:
                    existing_val = existing_row_data[col_idx] if col_idx < len(existing_row_data) else None
                    if existing_val is not None and existing_val != '':
                        continue
                    for row_idx in range(2, ws.max_row + 1):
                        if str(ws.cell(row=row_idx, column=key_idx + 1).value or '').strip().lower() == key:
                            ws.cell(row=row_idx, column=col_idx + 1, value=_to_excel_val(row_data[field]))
                            break
        else:
            row_values = [_to_excel_val(row_data.get(f, None)) for f in fields]
            ws.append(row_values)
            existing[key] = row_values

    wb.save(EXCEL_PATH)


class ExcelPipeline:
    def __init__(self):
        self.previred_items = []
        self.utm_items = []

    def open_spider(self, spider):
        _init_excel()

    def process_item(self, item, spider):
        if spider.name == 'previred':
            self.previred_items.append(dict(item))
        elif spider.name == 'sii_utm':
            self.utm_items.append(dict(item))
        return item

    def close_spider(self, spider):
        if spider.name == 'previred' and self.previred_items:
            _upsert_rows(self.previred_items, 'mes', PREVIRED_FIELDS)
        elif spider.name == 'sii_utm' and self.utm_items:
            _upsert_rows(self.utm_items, 'mes', UTM_FIELDS)
